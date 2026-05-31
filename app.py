import streamlit as st
import asyncio, subprocess, sys, os
import pandas as pd
from datetime import datetime, timedelta, date
from io import BytesIO
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


@st.cache_resource
def install_playwright():
    subprocess.run(["playwright", "install", "chromium", "--with-deps"],
                   capture_output=True)


install_playwright()
from playwright.async_api import async_playwright

URL = "https://tpb.vn/cong-cu-tinh-toan/ty-gia-ngoai-te"
KEEP_CURRENCIES = {"USD", "JPY", "AUD", "SGD"}


def to_excel_bytes(rows):
    df = pd.DataFrame(rows, columns=[
        "Ngay", "Ma ngoai te", "Ten ngoai te",
        "Mua chuyen khoan", "Ban chuyen khoan", "Nguon", "Ghi chu"
    ])
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Ty gia")
        ws = writer.book["Ty gia"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        fill = PatternFill("solid", fgColor="D9EAF7")
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for col_cells in ws.columns:
            w = max(len(str(c.value or "")) for c in col_cells)
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(w + 2, 60)
    buf.seek(0)
    return buf.read()


async def scrape_one_day(page, date_obj):
    date_text = date_obj.strftime("%d/%m/%Y")
    await page.reload(wait_until="domcontentloaded")
    await page.wait_for_timeout(3000)
    js_code = """
        (d) => {
            const picker = document.getElementById("datepickerInput");
            const from_el = document.getElementById("from_date");
            const to_el = document.getElementById("to_date");
            if (picker) picker.value = d;
            if (from_el) from_el.value = d;
            if (to_el) to_el.value = d;
            [picker, from_el, to_el].forEach(el => {
                if (!el) return;
                el.dispatchEvent(new Event("change", {bubbles: true}));
                el.dispatchEvent(new Event("input", {bubbles: true}));
            });
        }
    """
    await page.evaluate(js_code, date_text)
    await page.wait_for_timeout(500)
    js_sel = """
        () => {
            const selects = Array.from(document.querySelectorAll("select"));
            for (const s of selects) {
                if (s.options.length > 0) {
                    s.selectedIndex = s.options.length - 1;
                    s.dispatchEvent(new Event("change", {bubbles: true}));
                    return;
                }
            }
        }
    """
    await page.evaluate(js_sel)
    await page.wait_for_timeout(300)
    clicked = False
    for btn_sel in ["button:has-text('Xem ti gia')", "a:has-text('Xem ti gia')"]:
        try:
            if await page.locator(btn_sel).count() > 0:
                await page.locator(btn_sel).first.click(timeout=5000)
                clicked = True
                break
        except:
            pass
    await page.wait_for_timeout(4000)
    rows = []
    for tbl_sel in ["table tbody tr", "tbody tr"]:
        trs = page.locator(tbl_sel)
        count = await trs.count()
        if count > 0:
            for i in range(count):
                tds = trs.nth(i).locator("td")
                n = await tds.count()
                if n < 4:
                    continue
                vals = [(await tds.nth(j).inner_text()).strip() for j in range(n)]
                if vals[0].upper() not in KEEP_CURRENCIES:
                    continue
                rows.append({
                    "Ngay": date_text,
                    "Ma ngoai te": vals[0] if len(vals) > 0 else "",
                    "Ten ngoai te": vals[1] if len(vals) > 1 else "",
                    "Mua chuyen khoan": vals[3] if len(vals) > 3 else "",
                    "Ban chuyen khoan": vals[5] if len(vals) > 5 else "",
                    "Nguon": URL,
                    "Ghi chu": "",
                })
            break
    if not rows:
        rows.append({
            "Ngay": date_text, "Ma ngoai te": "", "Ten ngoai te": "",
            "Mua chuyen khoan": "", "Ban chuyen khoan": "",
            "Nguon": URL, "Ghi chu": "Khong doc duoc bang",
        })
    return rows


async def run_scraper(start_date, end_date, progress_cb, log_cb):
    all_rows = []
    total = (end_date - start_date).days + 1
    done = 0
    async with async_playwright() as p:
        browser = await p.chromium.launch(
            headless=True,
            args=["--no-sandbox", "--disable-dev-shm-usage",
                  "--disable-blink-features=AutomationControlled"]
        )
        ctx = await browser.new_context(
            viewport={"width": 1400, "height": 900},
            user_agent=("Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                        "AppleWebKit/537.36 Chrome/120.0.0.0 Safari/537.36")
        )
        page = await ctx.new_page()
        log_cb("Dang tai trang TPBank...")
        await page.goto(URL, wait_until="domcontentloaded", timeout=60000)
        await page.wait_for_timeout(5000)
        log_cb("Trang da tai xong")
        current = start_date
        while current <= end_date:
            try:
                rows = await scrape_one_day(page, current)
                all_rows.extend(rows)
                got = len([r for r in rows if r["Ma ngoai te"]])
                log_cb(f"  {current.strftime('%d/%m/%Y')} -> {got} dong")
            except Exception as e:
                log_cb(f"  {current.strftime('%d/%m/%Y')} -> Loi: {e}")
                all_rows.append({
                    "Ngay": current.strftime("%d/%m/%Y"),
                    "Ma ngoai te": "", "Ten ngoai te": "",
                    "Mua chuyen khoan": "", "Ban chuyen khoan": "",
                    "Nguon": URL, "Ghi chu": str(e),
                })
            done += 1
            progress_cb(done / total)
            current += timedelta(days=1)
        await browser.close()
    return all_rows


st.set_page_config(page_title="Ti gia TPBank", page_icon="💱", layout="centered")
st.title("💱 Ti gia ngoai te TPBank")
st.caption("Lay ti gia mua/ban chuyen khoan: USD · JPY · AUD · SGD")
st.divider()

col1, col2 = st.columns(2)
with col1:
    start = st.date_input("Tu ngay", value=date(2026, 3, 1),
                          min_value=date(2020, 1, 1), max_value=date(2030, 12, 31))
with col2:
    end = st.date_input("Den ngay", value=date(2026, 3, 5),
                        min_value=date(2020, 1, 1), max_value=date(2030, 12, 31))

st.divider()

if "result_bytes" not in st.session_state:
    st.session_state.result_bytes = None
if "log_lines" not in st.session_state:
    st.session_state.log_lines = []

if st.button("Lay ti gia", type="primary", use_container_width=True):
    if start > end:
        st.error("Ngay bat dau phai truoc ngay ket thuc!")
    else:
        st.session_state.result_bytes = None
        st.session_state.log_lines = []
        log_box = st.empty()
        prog_bar = st.progress(0, text="Dang khoi dong...")
        log_lines = []

        def log_cb(msg):
            log_lines.append(msg)
            log_box.code("\n".join(log_lines), language=None)

        def progress_cb(pct):
            prog_bar.progress(pct, text=f"{int(pct * 100)}% hoan thanh")

        start_dt = datetime.combine(start, datetime.min.time())
        end_dt = datetime.combine(end, datetime.min.time())
        rows = asyncio.run(run_scraper(start_dt, end_dt, progress_cb, log_cb))
        prog_bar.progress(1.0, text="Hoan thanh!")
        log_cb(f"Tong: {len(rows)} dong du lieu")
        st.session_state.result_bytes = to_excel_bytes(rows)
        st.session_state.log_lines = log_lines

if st.session_state.result_bytes:
    st.success("Du lieu da san sang!")
    fname = f"ty_gia_TPBank_{start}_{end}.xlsx"
    st.download_button(
        label="Tai file Excel",
        data=st.session_state.result_bytes,
        file_name=fname,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )
